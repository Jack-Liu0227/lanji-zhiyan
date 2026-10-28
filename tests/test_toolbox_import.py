import base64
import io
from pathlib import Path
import tempfile
import unittest

from docx import Document
from pypdf import PdfWriter
import workflow
import toolbox_import


class ToolboxImportTest(unittest.TestCase):
    def read(self,name,raw,tool='translate',db=None,**extra):
        return toolbox_import.import_file(db,{'tool':tool,'name':name,'data':base64.b64encode(raw).decode(),**extra})

    def test_text_encodings_and_supported_tools(self):
        for encoding in ('utf-8-sig','utf-16','gb18030'):
            self.assertEqual(self.read('论文.txt','研究内容'.encode(encoding))['text'],'研究内容')
        for tool,name in [('polish','paper.md'),('convert','paper.tex'),('references','refs.bib'),('mermaid','flow.mmd')]:
            self.assertEqual(self.read(name,b'CONTENT',tool)['text'],'CONTENT')
        with self.assertRaises(ValueError):self.read('figure.png',b'PNG')
        with self.assertRaises(ValueError):self.read('binary.txt',b'\x00binary')
        with self.assertRaises(ValueError):self.read('empty.txt',b' ')
        with self.assertRaises(ValueError):self.read('large.txt',b'a'*200001)

    def test_word_preserves_paragraph_table_order(self):
        document=Document();document.add_paragraph('First')
        table=document.add_table(rows=1,cols=2);table.cell(0,0).text='Cell A';table.cell(0,1).text='Cell B'
        document.add_paragraph('Last');buffer=io.BytesIO();document.save(buffer)
        self.assertEqual(self.read('paper.docx',buffer.getvalue())['text'],'First\n\nCell A\tCell B\n\nLast')

    def test_pdf_translation_import_and_validation(self):
        writer=PdfWriter();writer.add_blank_page(width=300,height=400);buffer=io.BytesIO();writer.write(buffer);raw=buffer.getvalue()
        with self.assertRaises(ValueError):self.read('bad.pdf',b'not a pdf')
        with self.assertRaisesRegex(ValueError,'没有可提取'):self.read('scan.pdf',raw)
        with tempfile.TemporaryDirectory() as folder, workflow.WorkflowDB(Path(folder)/'app.sqlite3') as db:
            project=Path(folder)/'project';project.mkdir();pid=db.create_task('Import',base_dir=str(project))
            with self.assertRaisesRegex(ValueError,'选择结果归属'):self.read('paper.pdf',raw,'pdf_translate',db)
            imported=self.read('paper.pdf',raw,'pdf_translate',db,project_id=pid)
            paper=db.paper(imported['paper_id'])
            self.assertEqual(paper['task_id'],pid)
            self.assertEqual(Path(paper['source_path']).read_bytes(),raw)
            self.assertTrue(Path(paper['source_path']).is_relative_to(project))
