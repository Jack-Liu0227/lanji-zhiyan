# -*- coding: utf-8 -*-
"""Markdown → Word（中文学位论文/期刊稿的常用排版规范）。

只做「排得像个稿子」这件事，不碰内容：
  A4、页边距 上下 2.54 / 左 3.0 / 右 2.5 cm
  正文  宋体 + Times New Roman 小四(12pt)、1.5 倍行距、首行缩进 2 字符、两端对齐
  题目  #  → 黑体 三号(16pt) 居中
  标题  ## → 黑体 三号 居中；### → 四号(14pt)；#### → 小四(12pt) 左顶格
  图表  图片居中宽 ≤15cm、图题五号居中；| 表格 | 转三线表（五号、居中）
  页脚  居中页码

python-docx 的默认模板是英文的（Letter 纸、Calibri、蓝色标题），直接 add_paragraph
出来在 Word 里就是「没排版」——所以这里全部按样式覆盖，不逐段硬写。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLACK = RGBColor(0, 0, 0)
SONG, HEI, LATIN = "宋体", "黑体", "Times New Roman"
BODY_PT, SMALL_PT = 12.0, 10.5
MAX_IMG_CM = 15.0


def _rfonts(el, ea: str, latin: str = LATIN) -> None:
    """把 rFonts 的 ascii/hAnsi/eastAsia 全部钉死，并清掉主题字体（否则 Word 按主题回填）。"""
    rPr = el.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rf)
    for k in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn("w:" + k), None)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), ea)


def _first_line(p, chars: int = 2, size: float = BODY_PT) -> None:
    """正文段落首行缩进 N 字符（挂段落级：挂样式会被标题/列表/页脚一起继承）。"""
    ind = p._p.get_or_add_pPr().get_or_add_ind()
    ind.set(qn("w:firstLineChars"), str(chars * 100))
    ind.set(qn("w:firstLine"), str(int(chars * size * 20)))


def _no_indent(p) -> None:
    """段落级取消首行缩进：firstLine 和 firstLineChars 必须一起清，只清一个会被样式里的 2 字符顶回来。"""
    p.paragraph_format.first_line_indent = Cm(0)
    ind = p._p.get_or_add_pPr().get_or_add_ind()
    ind.set(qn("w:firstLineChars"), "0")
    ind.set(qn("w:firstLine"), "0")


def _style(doc, name: str, ea: str, size: float, *, bold=False, align=None,
           first_line_chars=0, before=0, after=0, line=1.5) -> None:
    st = doc.styles[name]
    st.font.size, st.font.bold, st.font.italic = Pt(size), bold, False
    st.font.color.rgb = BLACK
    _rfonts(st.element, ea)
    pf = st.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    if align is not None:
        pf.alignment = align
    pPr = st.element.get_or_add_pPr()
    ind = pPr.get_or_add_ind()
    for k in ("firstLine", "firstLineChars", "hanging", "hangingChars"):
        ind.attrib.pop(qn("w:" + k), None)
    # 显式写 0，别留空让 Normal 的值继承下来
    ind.set(qn("w:firstLineChars"), str(first_line_chars * 100))
    ind.set(qn("w:firstLine"), str(int(first_line_chars * size * 20)))


def _setup(doc) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.5)
    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    C = WD_ALIGN_PARAGRAPH.CENTER
    _style(doc, "Normal", SONG, BODY_PT, align=J)
    _style(doc, "Title", HEI, 16, bold=True, align=C, after=18)
    _style(doc, "Heading 1", HEI, 16, bold=False, align=C, before=12, after=12)
    _style(doc, "Heading 2", HEI, 14, bold=False, before=6, after=6)
    _style(doc, "Heading 3", HEI, 12, bold=False, before=6, after=0)
    _style(doc, "Heading 4", HEI, 12, bold=True, before=6, after=0)
    for nm in ("List Bullet", "List Number"):
        _style(doc, nm, SONG, BODY_PT, after=0)
    _style(doc, "Footer", SONG, SMALL_PT, align=C)
    p = sec.footer.paragraphs[0]
    p.alignment = C
    _no_indent(p)
    fld = p._p.makeelement(qn("w:fldSimple"), {qn("w:instr"): " PAGE "})
    run = p.add_run()
    run.font.size = Pt(SMALL_PT)
    _rfonts(run._r, SONG)
    fld.append(run._r)
    p._p.append(fld)


def _inline(p, text: str) -> None:
    """只认 **粗体**；其余原样。"""
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not seg:
            continue
        r = p.add_run(seg)
        if i % 2:
            r.bold = True


def _three_line(table) -> None:
    """三线表：上下框 1.5pt、表头下框 0.75pt，无竖线。"""
    borders = table._tbl.tblPr.makeelement(qn("w:tblBorders"), {})
    for edge, sz in (("top", 12), ("left", 0), ("bottom", 12), ("right", 0),
                     ("insideH", 0), ("insideV", 0)):   # schema 顺序就是 top,left,bottom,right,insideH,insideV
        e = borders.makeelement(qn("w:" + edge), {})
        if sz:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:color"), "000000")
        else:
            e.set(qn("w:val"), "none")
            e.set(qn("w:sz"), "0")
        borders.append(e)
    # tblBorders 在 tblPr 里有固定位置，直接 append 会让 Word 报「内容有问题」
    table._tbl.tblPr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcB = tcPr.makeelement(qn("w:tcBorders"), {})
        tcB.append(tcB.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "6", qn("w:color"): "000000"}))
        tcPr.insert_element_before(tcB, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def build(markdown: str, dest: Path | str, bases=()) -> Path:
    """markdown → dest(.docx)。bases：解析图片相对路径的候选目录。"""
    doc = Document()
    _setup(doc)
    lines = (markdown or "").splitlines()
    seen_title = False
    i = 0
    while i < len(lines):
        line = lines[i]
        i += 1
        stripped = line.strip()
        if not stripped:
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            lvl, txt = len(m.group(1)), m.group(2).strip()
            if lvl == 1 and not seen_title:
                seen_title = True
                p = doc.add_paragraph(txt, style="Title")
            else:
                p = doc.add_paragraph(txt, style=f"Heading {min(max(lvl - 1, 1), 4)}")
            continue

        if stripped.startswith("|") and i < len(lines) and re.match(r"^\|[\s:|-]+\|?$", lines[i].strip()):
            head = _cells(stripped)
            i += 1                                             # 跳过 |---|---|
            body = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                body.append(_cells(lines[i]))
                i += 1
            t = doc.add_table(rows=1 + len(body), cols=len(head))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in ((t.rows[0], head), *((t.rows[n + 1], b) for n, b in enumerate(body))):
                for c, txt in enumerate(row[:len(head)]):
                    cell = r.cells[c]
                    cell.text = ""
                    par = cell.paragraphs[0]
                    par.style = doc.styles["Normal"]
                    _no_indent(par)
                    par.paragraph_format.line_spacing = 1.0
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = par.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", txt))
                    run.font.size = Pt(SMALL_PT)
                    _rfonts(run._r, SONG)
            _three_line(t)
            _no_indent(doc.add_paragraph("", style="Normal"))
            continue

        m = re.match(r"^!\[(.*?)\]\((.+?)\)\s*$", stripped)
        if m:
            cap, raw = m.group(1), m.group(2).strip("<>")
            path = next((q for q in ((Path(b) / raw if b else Path(raw)) for b in (bases or ("",))) if q.is_file()), None)
            if path:
                doc.add_picture(str(path))
                pic = doc.paragraphs[-1]
                if pic.runs and pic.runs[0].width and pic.runs[0].width > Cm(MAX_IMG_CM):
                    pic.runs[0].width, pic.runs[0].height = Cm(MAX_IMG_CM), None
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                _no_indent(doc.paragraphs[-1])
                if cap:
                    p = doc.add_paragraph(cap, style="Normal")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _no_indent(p)
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(12)
                    for r in p.runs:
                        r.font.size = Pt(SMALL_PT)
                        _rfonts(r._r, SONG)
                continue
            # 找不到图就当正文写出来，不放空（别静默丢东西）

        if stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            _inline(p, stripped[2:])
            continue

        p = doc.add_paragraph(style="Normal")
        _first_line(p)
        _inline(p, stripped)

    dest = Path(dest)
    doc.save(str(dest))
    return dest


if __name__ == "__main__":                                  # 自检：python workflow_docx.py out.docx
    import sys, tempfile
    md = ("# 题目 β-Ga₂O₃\n\n## 1. 引言\n\n### 1.1 背景\n\n首段正文，用于验证缩进。\n\n"
          "- 要点一\n- 要点二\n\n| 样品 | 暗电流 |\n|---|---|\n| A | 1e-9 |\n| B | 3e-9 |\n")
    out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(tempfile.gettempdir()) / "docx_check.docx")
    build(md, out)
    d = Document(str(out))
    s = d.sections[0]
    assert round(s.page_width.cm) == 21 and round(s.page_height.cm) == 30, "不是 A4"
    n = d.styles["Normal"]
    assert n.font.size == Pt(BODY_PT) and n.paragraph_format.line_spacing == 1.5
    body = next(x for x in d.paragraphs if x.text.startswith("首段正文"))
    assert body._p.pPr.ind.get(qn("w:firstLineChars")) == "200", "首行缩进不是 2 字符"
    assert n.element.rPr.rFonts.get(qn("w:eastAsia")) == SONG
    assert d.styles["Heading 1"].element.rPr.rFonts.get(qn("w:eastAsia")) == HEI
    assert d.paragraphs[0].style.name == "Title"
    assert d.tables and len(d.tables[0].rows) == 3, "表格没转出来"
    assert d.sections[0].footer.paragraphs[0]._p.findall(qn("w:fldSimple")), "页脚没有页码域"
    print("ok →", out)
