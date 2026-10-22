# -*- coding: utf-8 -*-
"""七步工作流阶段逻辑（创新点审查/检索入库/精读/知识库笔记/撰写审查）模块。"""
from __future__ import annotations

import re
import threading
from pathlib import Path
from typing import TYPE_CHECKING, Any

import paths
from workflow_prompts import READ_SYSTEM
from workflow_search import MAX_RESULTS, download_pdf, enrich_oa, query_tokens, rank, search

if TYPE_CHECKING:
    from workflow import WorkflowDB

PAPERS = paths.PAPERS
NOTES = paths.NOTES
UA = {"User-Agent": "research-workbench-1.0 (mailto:local@localhost)"}

def _text(value: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", value or "")).strip()

def stage_idea(db: WorkflowDB, task_id: str, idea: str) -> dict[str, Any]:
    from workflow import chat_json
    task = db.summary(task_id)["task"]
    profile = db.active_model()
    out = chat_json(profile, [
        {"role": "system", "content": '你是严格的科研选题评审（博导视角）。只返回 JSON：'
         '{"type":"Novel Problem|Novel Method|New Setting","story":"一句话讲清","flaws":[{"flaw":"","severity":"CRITICAL|MAJOR","defense":""}],'
         '"dims":{"Higher":0-10,"Faster":0-10,"Stronger":0-10,"Cheaper":0-10,"Broader":0-10},'
         '"dims_f":{"F1":0-10,"F2":0-10,"F3":0-10,"F4":0-10,"F5":0-10},'
         '"dims_evidence":{"F1":"","F2":"","F3":"","F4":"","F5":""},'
         '"falsifications":[{"risk":"HIGH|MEDIUM","criticism":"审稿人潜在质疑","defense":"防御实验设计"}],'
         '"verdict":"Strong Accept|Accept with Revisions|Reject and Pivot","reason":"","next":["下一步 1","下一步 2"]}'
         '。F1=理论突破度，F2=逻辑闭环度，F3=实验完备性，F4=数据可达性，F5=领域紧迫性。'
         '反向证伪必须提出3种审稿人最尖锐的质疑与Reject风险，并设计防御实验。判 Strong Accept 必须 ≥2 个维度 ≥8 且无 CRITICAL 缺陷。'},
        {"role": "user", "content": f"研究任务：{task['title']}\n研究问题：{task['question']}\n关键词 / 材料体系：{task['material_system']}\n\n待审查创新点：\n{idea}"},
    ])
    dims = out.get("dims") or {}
    dim_txt = "　".join(f"{k}：{v}/10" for k, v in dims.items() if v is not None) or "（未给分）"
    f_dims = out.get("dims_f") or {}
    f_labels = {"F1": "F1 理论突破度", "F2": "F2 逻辑闭环度", "F3": "F3 实验完备性", "F4": "F4 数据可达性", "F5": "F5 领域紧迫性"}
    f_txt = "　".join(f"{f_labels.get(k, k)}：{v}/10" for k, v in f_dims.items() if v is not None) if f_dims else ""
    flaws = out.get("flaws") or []
    flaw_lines = [f"- [{x.get('severity')}] {x.get('flaw')}　→ {x.get('defense')}" for x in flaws] or ["- 无"]
    fals = out.get("falsifications") or []
    fal_lines = [f"- **[{x.get('risk','MEDIUM')}] 质疑**：{x.get('criticism','')}\n  - 🛡️ **防御实验**：{x.get('defense','')}" for x in fals] or []
    next_lines = [f"- {x}" for x in (out.get("next") or [])]
    md_parts = [
        f"# 创新点审查 · {task['title']}", "",
        f"- 类型：{out.get('type','')}", f"- 一句话：{out.get('story','')}", "",
        f"- 五维指标：{dim_txt}",
    ]
    if f_txt:
        md_parts.append(f"- 博导评审矩阵：{f_txt}")
    md_parts += [
        f"- 判定：**{out.get('verdict','')}**", "", f"**理由**：{out.get('reason','')}", "",
        "## 致命缺陷", *flaw_lines,
    ]
    if fal_lines:
        md_parts += ["", "## 反向证伪与防御实验设计", *fal_lines]
    md_parts += [
        "", "## 下一步", *next_lines, "",
        "> 模型意见是草稿式审查，最终判断以导师与证据为准。",
    ]
    md = "\n".join(md_parts)
    note = db.add_note(task_id, f"创新点审查 · {task['title']}", md, category="研究计划")
    db.log_step(task_id, "idea", "done", f"判定 {out.get('verdict','')}")
    db.add_memory(task_id, "long", f"创新点审查结论：{out.get('verdict','')}｜理由：{out.get('reason','')[:160]}")
    return {"verdict": out, "note_id": note["id"], "markdown": md}


def _relevant(query: str, item: dict[str, Any]) -> bool:
    """相关度粗筛：查询里 ≥2 个实词（≥4 字符）出现在标题+摘要里才算命中。
    目的：挡掉「Teacher materials Volume 1」这类只蹭到一个词的噪声。关键词本来就少时不过滤。"""
    toks = query_tokens(query)
    if len(toks) < 2:
        return True
    blob = ("%s %s" % (item.get("title") or "", item.get("abstract") or "")).lower()
    return sum(1 for t in toks if t in blob) >= 2


def stage_search(db: WorkflowDB, task_id: str, query: str, sources: list[str] | None = None, rows: int = 30,
                 limit: int = MAX_RESULTS) -> dict[str, Any]:
    query = (query or "").strip() or db.summary(task_id)["task"]["question"]
    if not query:
        raise ValueError("请填写检索关键词")
    rows = max(1, min(int(rows or 30), 50))       # 单源上限 50，三源够凑满 limit
    limit = max(1, min(int(limit or MAX_RESULTS), MAX_RESULTS))
    raw = search(query, sources, rows)
    rel = [x for x in raw if _relevant(query, x)]
    found = rank(query, rel, limit)            # 去重 + 按 相关度×引用量 排序 + 限条数（在补 OA 前截，少发请求）
    filled = enrich_oa(found)                  # 按 DOI 反查 OpenAlex，尽量把 OA 直链补全
    saved = db.cache_results(task_id, found)
    db.log_step(task_id, "search", "done",
                f"检索到 {len(found)} 条（原始 {len(raw)}，粗筛掉 {len(raw) - len(rel)}，去重排序后限 {limit} 条），"
                f"新增 {len(saved)} 条，补 OA 直链 {filled} 条")
    return {"query": query, "found": len(found), "raw": len(raw), "filtered": len(raw) - len(rel),
            "limit": limit, "saved": saved, "enriched": filled, "results": db.results(task_id)}


def stage_download(db: WorkflowDB, task_id: str, result_id: str) -> dict[str, Any]:
    if not (result_id or "").strip():
        raise ValueError("没给要下载的那条检索结果 —— 在「论文搜索」里点某条的「下载 OA」，或勾选后自动下载")
    item = db.result(result_id)
    if not item:
        raise ValueError("这条检索结果不在了（可能被清过），重新搜一次")
    path = download_pdf(item, target_dir=db.pdir(task_id, "文献库"))
    paper = db.ingest_pdf(task_id, path)
    db.log_step(task_id, "download", "done", f"已入库 {paper['filename']}（{len(paper['pages'])} 页）")
    return {"paper": paper, "pdf_path": str(path)}


def stage_download_selected(db: WorkflowDB, task_id: str, result_ids: list[str]) -> dict[str, Any]:
    """论文搜索勾选 → 处理：能自动下的下到「文献库」并入库；下不到（无 OA / 403 / 不是 PDF）的进「待手动下载清单」。"""
    ids = [i for i in (result_ids or []) if i]
    if not ids:
        raise ValueError("先在「论文搜索」里勾选论文")
    db.enqueue(task_id, ids)
    done = manual = 0
    for r in db.queue(task_id):
        item = db.result(r["id"])
        if not (item.get("pdf_url") or "").strip():
            db.set_queue(task_id, r["id"], "manual", "", "无开放获取(OA) 全文，需手动下载")
            manual += 1
            continue
        try:
            path = download_pdf(item, target_dir=db.pdir(task_id, "文献库"))
            db.ingest_pdf(task_id, path)
            db.set_queue(task_id, r["id"], "downloaded", str(path), "")
            done += 1
        except Exception as exc:                       # 403 / 落地页不是 PDF / 超时 → 交给人工
            db.set_queue(task_id, r["id"], "manual", "", f"自动下载失败：{exc}")
            manual += 1
    db.log_step(task_id, "download", "done", f"勾选 {len(ids)} 条：自动下载 {done} 篇，待手动 {manual} 篇")
    return {"selected": len(ids), "downloaded": done, "manual": manual, "items": db.queue(task_id)}


def stage_zotero_import(db: WorkflowDB, task_id: str, collection: str = "",
                        keys: list[str] | None = None,
                        collections: list[str] | None = None) -> dict[str, Any]:
    """Zotero 分类（支持单个或多个 collection 接管，或点选的几条）→ 直接进「下载精读」队列。

    题录用 Zotero 的（零错误、不花 LLM），PDF 用本机附件直读，不检索、不联网下载、不复制文件。
    Zotero 里没 PDF 附件的条目照样进队列，标成「待手动」，不假装成功。
    """
    import zotero                                     # 只走本机 Local API；没启动会抛中文原因
    keys = [k for k in (keys or []) if k]
    if keys:
        rows = [zotero.item(k) for k in keys]
    else:
        coll_list = []
        if collections:
            coll_list = [c.strip() for c in collections if c.strip()]
        elif collection:
            coll_list = [c.strip() for c in collection.split(",") if c.strip()]
        if not coll_list:
            coll_list = [""]

        rows, seen_keys = [], set()
        for col in coll_list:
            start = 0
            while True:
                page = zotero.items(collection=col, start=start, limit=zotero.MAX_LIMIT)
                for r in page["items"]:
                    if r["type"] not in ("attachment", "note", "annotation") and r["key"] not in seen_keys:
                        seen_keys.add(r["key"])
                        rows.append(r)
                start += page["returned"]
                if not page["hasMore"] or not page["returned"]:
                    break
    done = manual = 0
    for it in rows:
        rid = db.cache_result(task_id, {
            "source": "zotero", "title": it["title"], "authors": ", ".join(it["creators"]),
            "year": it["year"], "doi": it["doi"],
            "url": it["url"] or (f"https://doi.org/{it['doi']}" if it["doi"] else ""),
            "pdf_url": "", "abstract": it["abstract"],
        })
        path = zotero.pdf_path(it["key"])
        if not path:
            db.set_queue(task_id, rid, "manual", "", "Zotero 里这条没有 PDF 附件")
            manual += 1
            continue
        try:
            db.ingest_pdf(task_id, path)
        except Exception as exc:                      # 加密/损坏/不是真 PDF → 交给人工
            db.set_queue(task_id, rid, "manual", "", f"Zotero 附件读不了：{exc}")
            manual += 1
            continue
        db.set_queue(task_id, rid, "downloaded", path, "")
        done += 1
    if done:
        try:
            db.build_paper_graph(task_id)
            db.kb_graph(task_id, rebuild=True)
        except Exception:
            pass
    db.log_step(task_id, "download", "done",
                f"Zotero 导入 {len(rows)} 条：入库 {done} 篇，无附件 {manual} 条")
    return {"total": len(rows), "downloaded": done, "manual": manual, "items": db.queue(task_id)}


def is_paper_pdf(path: Any) -> bool:
    """这算不算「一篇论文」：overleaf 编出来的成品、_build 里的、旁边有同名 .tex 的、<20KB 的都跳过。"""
    p = Path(str(path))
    low = str(p).replace("\\", "/").lower()
    if "_build" in low or "/overleaf/" in low or "/__pycache__/" in low:
        return False
    try:
        if p.stat().st_size < 20000:          # 太小的多半是封面/编译残件，不是论文
            return False
    except OSError:
        return False
    if p.with_suffix(".tex").exists():        # 旁边就是它的 .tex 源 → 这是编出来的
        return False
    if ".no_watermark.zh." in low:            # BabelDOC 译文（xxx.no_watermark.zh.mono|dual.pdf）不是新论文
        return False
    return True


class DeepreadStopped(Exception):
    """用户点了停止 —— 不是失败，是主动中断。"""


DEEPREAD: dict[str, dict[str, Any]] = {}      # task_id -> 进度（内存；重启即丢，这是刻意的）


def deepread_state(task_id: str) -> dict[str, Any] | None:
    return DEEPREAD.get(task_id)


def deepread_stop(task_id: str) -> dict[str, Any]:
    st = DEEPREAD.get(task_id)
    if st:
        st["stop"] = True
    return {"ok": True, "stopping": bool(st)}


def deepread_kick(db_path: str, task_id: str) -> dict[str, Any]:
    """把「一键精读」扔到后台跑（原来同步跑，篇多会卡死界面、也停不下来）。"""
    if task_id in DEEPREAD:
        return {"started": False, "reason": "这个项目已经在精读了"}
    DEEPREAD[task_id] = {"total": 0, "done": 0, "failed": 0, "skipped": [],
                         "current": "准备中", "phase": "running", "stop": False, "errors": []}
    threading.Thread(target=_deepread_worker, args=(db_path, task_id), daemon=True).start()
    return {"started": True}


READ_NOTE_TITLE = "精读 · "
MIN_READ_NOTE_CHARS = 300     # 只有标题的空笔记不算「读过了」（以前模型回空留下过这种）


def read_done_filenames(db: Any, task_id: str) -> set[str]:
    """已有实质内容的精读笔记覆盖的 PDF 文件名 —— 重跑时不重读，别重复烧 token。"""
    done = set()
    for n in db.notes(task_id):
        title = n["title"] or ""
        if title.startswith(READ_NOTE_TITLE) and len((n["markdown"] or "").strip()) >= MIN_READ_NOTE_CHARS:
            done.add(title[len(READ_NOTE_TITLE):])
    return done


def _deepread_worker(db_path: str, task_id: str) -> None:
    st = DEEPREAD[task_id]
    try:
        with WorkflowDB(db_path) as db:
            folder = db.pdir(task_id, "文献库")
            known = {p["filename"] for p in db.papers_of(task_id)}
            skipped = []
            for pdf in sorted(folder.glob("*.pdf")):
                if st.get("stop"):
                    st["phase"] = "stopped"
                    return
                if pdf.name in known:
                    continue
                if not is_paper_pdf(pdf):
                    skipped.append(pdf.name)
                    continue
                st["current"] = "入库 " + pdf.name[:44]
                try:
                    db.ingest_pdf(task_id, pdf)
                except Exception as exc:
                    st["failed"] += 1
                    st["errors"].append("入库 %s: %s" % (pdf.name[:30], exc))
            st["skipped"] = skipped
            papers = db.papers_of(task_id)
            if not papers:
                raise ValueError("这个项目的「文献库」里没有可精读的 PDF —— 先搜一篇下下来，或把 PDF 丢进文献库")
            read_ok = read_done_filenames(db, task_id)
            st["total"] = len(papers)
            st["done"] = sum(1 for p in papers if p["filename"] in read_ok)      # 已精读过的算完成，不重读
            for p in papers:
                if st.get("stop"):
                    st["phase"] = "stopped"
                    return
                if p["filename"] in read_ok:
                    continue                       # 已经有实质内容的精读笔记，重跑跳过
                st["current"] = p["filename"][:46]
                try:
                    stage_read(db, task_id, p["id"], should_stop=lambda: bool(st.get("stop")))
                except DeepreadStopped:
                    st["phase"] = "stopped"
                    return
                except Exception as exc:
                    st["failed"] += 1
                    st["errors"].append("%s: %s" % (p["filename"][:30], exc))
                st["done"] += 1
            db.log_step(task_id, "notes", "done", "精读 %d 篇（失败 %d，跳过非论文 %d）" % (st["done"], st["failed"], len(skipped)))
            st["phase"] = "done"
    except Exception as exc:
        st["phase"] = "failed"
        st["error"] = "%s: %s" % (type(exc).__name__, exc)
    finally:
        st["current"] = ""




def stage_read(db: WorkflowDB, task_id: str, paper_id: str, per_page: int = 6000,
               should_stop: Any = None) -> dict[str, Any]:
    """逐页精读：把每页文本连同页码交给模型，要求中文要点 + 可引用原文。"""
    from workflow import chat
    profile = db.active_model()
    paper = db.paper(paper_id)
    if paper["text_chars"] < 100:
        raise ValueError("这篇 PDF 疑似扫描件，没有文本层，无法精读。")
    chunks, buf, size = [], [], 0
    for p in paper["pages"]:
        block = f"[PAGE {p['page_no']}]\n{p['text'][:per_page]}"
        if size + len(block) > 24000 and buf:
            chunks.append("\n\n".join(buf)); buf, size = [], 0
        buf.append(block); size += len(block)
    if buf:
        chunks.append("\n\n".join(buf))
    parts = []
    for i, chunk in enumerate(chunks, 1):
        if should_stop and should_stop():      # 每段检查一次：停就立刻不读了
            raise DeepreadStopped()
        text = chat(profile, [
            {"role": "system", "content": READ_SYSTEM},
            {"role": "user", "content": f"论文：{paper['filename']}（第 {i}/{len(chunks)} 段）\n\n{chunk}"},
        ], max_tokens=6000).strip()
        if not text:
            # 空回就别落一条只有标题的空笔记 —— 那会让「精读完了」变成假进度
            raise RuntimeError(f"第 {i}/{len(chunks)} 段没有拿到内容，模型回空，这篇先算失败")
        parts.append(text)
    md = f"# 精读 · {paper['filename']}\n\n" + "\n\n".join(parts)
    note = db.add_note(task_id, READ_NOTE_TITLE + paper["filename"], md, paper_id)
    db.log_step(task_id, "download", "done", f"精读完成：{paper['filename']}")
    return {"note_id": note["id"], "markdown": md, "chunks": len(chunks)}


def stage_notes(db: WorkflowDB, task_id: str, note_id: str, vault: str | None = None, apply: bool = False) -> dict[str, Any]:
    """把精读笔记整理成结构化知识库笔记（含证据表），可选写入 Obsidian。"""
    src = db.note(note_id)
    rows = [dict(r) for r in db.evidence(task_id)]
    prof = db.active_model()
    md = src["markdown"]
    # 先用严格闸门抽取候选证据，拒绝页码/原文对不上的条目
    from research_core import extract_ai
    if src.get("paper_id") and prof:
        try:
            extract_ai(db, task_id, src["paper_id"], prof["base_url"], prof["api_key"], prof["model"])
        except Exception:
            pass  # 抽不到证据不阻塞笔记整理
    rows = [dict(r) for r in db.evidence(task_id)]
    lines = [f"# 知识库笔记 · {src['title']}", "", "## 要点", "", md, "", "## 证据（回查原文）", ""]
    lines += [f"- **{r['field']}**：{r['value']}（{r['filename']}，p.{r['page_no']}）\n  - 原文：{r['excerpt']}" for r in rows] or ["- 暂无"]
    body = "\n".join(lines)
    note = db.add_note(task_id, f"笔记 · {src['title']}", body, src.get("paper_id"))
    result = {"note_id": note["id"], "markdown": body, "evidence": len(rows)}
    if vault:
        from integrations import apply_obsidian_preview, obsidian_preview
        preview = obsidian_preview(vault, src["title"], body)
        result["preview"] = {k: preview[k] for k in ("target", "exists", "confirm_token")}
        if apply:
            result["written"] = apply_obsidian_preview(preview, preview["confirm_token"])
    db.log_step(task_id, "notes", "done", f"知识库笔记已生成（证据 {len(rows)} 条）")
    return result


def stage_write(db: WorkflowDB, task_id: str, title: str = "", section: str = "引言") -> dict[str, Any]:
    from workflow import chat
    profile = db.active_model()
    task = db.summary(task_id)["task"]
    notes = db.notes(task_id)
    ev = [dict(r) for r in db.evidence(task_id)]
    context = "\n\n".join(f"### {n['title']}\n{n['markdown'][:6000]}" for n in notes[:6])
    ev_text = "\n".join(f"- {r['field']}: {r['value']}（{r['filename']} p.{r['page_no']}）" for r in ev[:60])
    md = chat(profile, [
        {"role": "system", "content": f"你是中文材料科学论文写作助手。按《{section}》写，只使用给定资料，每个关键结论后用（文献名 p.页码）标注出处；资料没给的写「资料不足」，禁止编造数据或引用。"},
        {"role": "user", "content": f"题目：{title or task['title']}\n研究问题：{task['question']}\n关键词：{task['material_system']}\n\n== 笔记 ==\n{context}\n\n== 证据 ==\n{ev_text}"},
    ], max_tokens=4000)
    doc = f"# {title or task['title']} · {section}\n\n{md}\n"
    out = db.pdir(task_id, "论文正文") / f"draft-{task_id[:8]}-{section}.md"
    out.write_text(doc, encoding="utf-8")
    db.log_step(task_id, "write", "done", f"{section} 草稿 {len(md)} 字")
    return {"markdown": doc, "path": str(out)}


def stage_review(db: WorkflowDB, task_id: str, draft: str) -> dict[str, Any]:
    from workflow import chat_json
    profile = db.active_model()
    ev = [dict(r) for r in db.evidence(task_id)]
    ev_text = "\n".join(f"- {r['field']}: {r['value']}（{r['filename']} p.{r['page_no']}）" for r in ev[:60])
    out = chat_json(profile, [
        {"role": "system", "content": "你是严格审稿人。只返回 JSON：{\"verdict\":\"接受|小修|大修|拒稿\",\"scores\":{\"证据\":0-5,\"逻辑\":0-5,\"表达\":0-5},\"issues\":[{\"severity\":\"高|中|低\",\"quote\":\"原文片段\",\"problem\":\"\",\"fix\":\"\"}],\"missing_evidence\":[\"...\"]}"},
        {"role": "user", "content": f"== 稿件 ==\n{draft[:24000]}\n\n== 已登记证据 ==\n{ev_text}"},
    ])
    md = "\n".join([
        "# 稿件审查", "", f"- 结论：**{out.get('verdict','')}**", f"- 评分：{out.get('scores',{})}", "",
        "## 问题", *[f"- [{i.get('severity')}] {i.get('problem')}\n  - 原文：{i.get('quote')}\n  - 建议：{i.get('fix')}" for i in out.get("issues", [])], "",
        "## 证据缺口", *[f"- {x}" for x in out.get("missing_evidence", [])], "",
        "> 审查意见为草稿，是否采纳由作者判断。",
    ])
    note = db.add_note(task_id, "稿件审查", md, category="写作投稿")
    db.log_step(task_id, "review", "done", f"结论 {out.get('verdict','')}")
    return {"review": out, "markdown": md, "note_id": note["id"]}


def stage_convert(db: WorkflowDB, task_id: str, markdown: str, fmt: str = "docx", name: str = "paper") -> dict[str, Any]:
    fmt = (fmt or "docx").lower()
    safe = re.sub(r"[^\w\-. ]+", "_", name)[:60] or "paper"
    folder = db.pdir(task_id, "最终产物")
    if fmt == "md":
        dest = folder / f"{safe}.md"; dest.write_text(markdown, encoding="utf-8")
    elif fmt == "html":
        import html as _html
        body = "\n".join(
            f"<h{min(len(m.group(1)), 6)}>{_html.escape(m.group(2))}</h{min(len(m.group(1)), 6)}>"
            if (m := re.match(r"^(#{1,6})\s+(.*)$", line))
            else (f"<li>{_html.escape(line[2:])}</li>" if line.startswith("- ") else f"<p>{_html.escape(line)}</p>" if line.strip() else "")
            for line in markdown.splitlines())
        dest = folder / f"{safe}.html"
        dest.write_text(f"<!doctype html><meta charset='utf-8'><body>{body}</body>", encoding="utf-8")
    elif fmt == "docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("未安装 python-docx，无法导出 docx（可改用 md/pdf 或 `py -3.10 -m pip install python-docx`）") from exc
        doc = Document()
        for line in markdown.splitlines():
            m = re.match(r"^(#{1,6})\s+(.*)$", line)
            if m:
                doc.add_heading(m.group(2), level=min(len(m.group(1)), 4))
            elif line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line.strip():
                doc.add_paragraph(line)
        dest = folder / f"{safe}.docx"; doc.save(str(dest))
    else:
        raise ValueError(f"不支持的格式：{fmt}（可用 md/html/docx）")
    db.log_step(task_id, "convert", "done", f"导出 {dest.name}")
    return {"path": str(dest), "format": fmt, "size": dest.stat().st_size}
